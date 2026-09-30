#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Scanly Vision Algo V2 architecture Word specification with diagrams."""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "doc" / "02-dr" / "视觉算法"
FIG_DIR = DOC_DIR / "figures"
OUT = DOC_DIR / "板材封边缺陷视觉算法架构说明书V2.docx"

# industrial palette
C_BG = "#F5F7FA"
C_EDGE = "#1F4E79"
C_EDGE_FILL = "#D6E3F0"
C_ALGO = "#0D7377"
C_ALGO_FILL = "#D5F0F1"
C_HW = "#5B5B5B"
C_HW_FILL = "#E8E8E8"
C_PLC = "#8B5A00"
C_PLC_FILL = "#F5E6C8"
C_ARROW = "#333333"
C_TEXT = "#1A1A1A"
C_ACCENT = "#C45C26"


def _setup_cn_font():
    plt.rcParams["font.sans-serif"] = [
        "PingFang SC",
        "Heiti SC",
        "STHeiti",
        "Songti SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def _box(ax, x, y, w, h, text, face, edge, fontsize=9, lw=1.4):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.02,rounding_size=0.08",
        facecolor=face,
        edgecolor=edge,
        linewidth=lw,
        mutation_aspect=0.3,
    )
    ax.add_patch(patch)
    ax.text(
        x + w / 2,
        y + h / 2,
        text,
        ha="center",
        va="center",
        fontsize=fontsize,
        color=C_TEXT,
        linespacing=1.25,
        wrap=True,
    )
    return patch


def _arrow(ax, x1, y1, x2, y2, text="", color=C_ARROW):
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="-|>", color=color, lw=1.3, mutation_scale=12),
    )
    if text:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        ax.text(mx, my + 0.08, text, ha="center", va="bottom", fontsize=7.5, color=color)


def _subtitle(ax, text, x=0.5, y=0.97):
    ax.text(x, y, text, transform=ax.transAxes, ha="center", va="top", fontsize=11, fontweight="bold", color=C_EDGE)


def draw_system_architecture(path: Path):
    """Overall system: field devices → edge client → algo V2 → upper systems."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 7.2), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    _subtitle(ax, "图1  系统总体架构")

    # left hardware
    _box(ax, 0.3, 6.2, 2.2, 0.9, "光电 / 编码器", C_HW_FILL, C_HW, 9)
    _box(ax, 0.3, 4.8, 2.2, 0.9, "PLC\n硬触发 TRIG", C_PLC_FILL, C_PLC, 9)
    _box(ax, 0.3, 3.0, 2.2, 1.4, "三路线扫相机\n上 / 中 / 下\nGigE · Mono8", C_HW_FILL, C_HW, 9)
    _arrow(ax, 1.4, 6.2, 1.4, 5.7)
    _arrow(ax, 1.4, 4.8, 1.4, 4.4)

    # edge existing group
    group = FancyBboxPatch(
        (3.0, 2.2),
        4.0,
        5.2,
        boxstyle="round,pad=0.03,rounding_size=0.1",
        facecolor="#F0F4F8",
        edgecolor=C_EDGE,
        linewidth=1.6,
        linestyle="-",
    )
    ax.add_patch(group)
    ax.text(5.0, 7.15, "边端客户端（现有，保留）", ha="center", va="center", fontsize=10, fontweight="bold", color=C_EDGE)
    _box(ax, 3.25, 5.7, 3.5, 0.95, "工程方案 / ROI / 阈值\nUI 与配方管理", C_EDGE_FILL, C_EDGE, 8.5)
    _box(ax, 3.25, 4.3, 3.5, 1.1, "采图模块\nSDK 曝光 · 硬触发取流 · 落盘", C_EDGE_FILL, C_EDGE, 8.5)
    _box(ax, 3.25, 2.55, 3.5, 1.4, "通讯与业务\n→ PLC 控制\n→ 视觉中心 / MES", C_EDGE_FILL, C_EDGE, 8.5)
    _arrow(ax, 5.0, 5.7, 5.0, 5.4)

    # camera to grab
    _arrow(ax, 2.5, 3.7, 3.25, 4.85, "图像")
    _arrow(ax, 2.5, 5.25, 3.25, 4.85, "TRIG 配置")

    # algo V2
    group2 = FancyBboxPatch(
        (7.5, 2.2),
        4.2,
        5.2,
        boxstyle="round,pad=0.03,rounding_size=0.1",
        facecolor="#EAF7F7",
        edgecolor=C_ALGO,
        linewidth=1.6,
    )
    ax.add_patch(group2)
    ax.text(9.6, 7.15, "视觉算法服务 V2（本机新增）", ha="center", va="center", fontsize=10, fontweight="bold", color=C_ALGO)

    steps = [
        (6.35, "帧接入_读本地文件"),
        (5.55, "预处理 / ROI / 切片"),
        (4.75, "检测 YOLO11s"),
        (3.95, "分类 YOLO11n-cls"),
        (3.15, "度量 + 单路判定"),
        (2.35, "板尾三路 OR 汇总"),
    ]
    for y, t in steps:
        _box(ax, 7.75, y, 3.7, 0.65, t, C_ALGO_FILL, C_ALGO, 8.5)
    for i in range(len(steps) - 1):
        y1 = steps[i][0]
        y2 = steps[i + 1][0] + 0.65
        _arrow(ax, 9.6, y1, 9.6, y2)

    # cross arrows
    _arrow(ax, 6.75, 4.85, 7.75, 6.5, "文件路径")
    _arrow(ax, 6.75, 6.15, 7.75, 5.85, "配方参数")
    _arrow(ax, 7.75, 2.65, 6.75, 3.25, "OK/NG\n缺陷列表")

    # legend note
    ax.text(
        6.0,
        0.35,
        "部署：边端同机 · 图像硬盘落盘+本地路径 · 算法优先Ubuntu(可沿用原OS) · 先原硬件复原再评估降本 · 不改中心/MES协议",
        ha="center",
        va="center",
        fontsize=7.5,
        color="#555555",
        style="italic",
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_module_boundary(path: Path):
    """Responsibility swimlane-style boundary diagram."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 7.2), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8.0)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图2  算法 V2 与周边软件模块边界")

    # zone banners (no overlap with body)
    zones = [
        (0.2, 4.4, "现场既有 · 电气", C_PLC_FILL, C_PLC),
        (2.5, 6.7, "边端软件（采图 + 客户端 + 算法 V2）", C_EDGE_FILL, C_EDGE),
        (9.4, 2.4, "产线上层", C_HW_FILL, C_HW),
    ]
    for x, w, title, face, edge in zones:
        _box(ax, x, 6.95, w, 0.45, title, face, edge, 8, lw=1.2)

    lanes = [
        (0.2, "电气 / PLC", C_PLC_FILL, C_PLC, "硬触发脉冲\n板材到位信号\n接收 NG 控制"),
        (2.5, "边端采图", C_EDGE_FILL, C_EDGE, "SDK 曝光参数\n硬触发取流\n图像落盘"),
        (4.8, "边端客户端", C_EDGE_FILL, C_EDGE, "工程方案/ROI\n下发文件路径\n结果展示与转发"),
        (7.1, "算法 V2", C_ALGO_FILL, C_ALGO, "按路径读图\n检测/分类/度量\n单路判定与汇总"),
        (9.4, "中心 / MES", C_HW_FILL, C_HW, "多机汇总复检\n业务数据\n协议不变"),
    ]
    for x, title, face, edge, body in lanes:
        _box(ax, x, 5.85, 2.1, 0.7, title, face, edge, 9, lw=1.6)
        _box(ax, x, 3.0, 2.1, 2.55, body, "#FFFFFF", edge, 8.5)

    # flow arrows between body columns
    for x in [2.3, 4.6, 6.9, 9.2]:
        _arrow(ax, x, 4.25, x + 0.2, 4.25)

    ax.text(6.0, 2.25, "数据方向（示意）", ha="center", fontsize=9, fontweight="bold", color=C_TEXT)
    flow_items = [
        "TRIG / 电气信号",
        "落盘图像文件",
        "路径 + 配方参数",
        "OK/NG + 缺陷",
        "原协议上报",
    ]
    xs = [1.25, 3.55, 5.85, 8.15, 10.45]
    for x, t in zip(xs, flow_items):
        _box(ax, x - 0.9, 1.05, 1.8, 0.75, t, C_BG, "#888888", 7.5)
    for i in range(len(xs) - 1):
        _arrow(ax, xs[i] + 0.9, 1.4, xs[i + 1] - 0.9, 1.4)

    # legend note at bottom — full width box, no overlap
    _box(
        ax,
        0.3,
        0.15,
        11.4,
        0.65,
        "说明：左侧电气；中间边端（采图落盘、客户端、算法V2按路径读图）；右侧中心/MES；算法OS优先Ubuntu，迁移动成本高可沿用原系统",
        "#FFF8F0",
        C_ACCENT,
        8,
        lw=1.0,
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_acquisition_flow(path: Path):
    """Hard trigger acquisition sequence."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 4.8), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图3  采图与硬触发协作流程")

    nodes = [
        (0.3, "光电检测到板"),
        (2.5, "PLC 输出 TRIG"),
        (4.7, "相机按设定曝光"),
        (6.9, "边端 SDK 取流并落盘"),
        (9.1, "V2 按路径读图\n帧接入"),
    ]
    y = 2.4
    w, h = 1.9, 1.15
    for i, (x, t) in enumerate(nodes):
        face = C_ALGO_FILL if i == 4 else (C_PLC_FILL if i == 1 else C_EDGE_FILL if i >= 3 else C_HW_FILL)
        edge = C_ALGO if i == 4 else (C_PLC if i == 1 else C_EDGE if i >= 3 else C_HW)
        _box(ax, x, y, w, h, t, face, edge, 8.5)
        if i < len(nodes) - 1:
            _arrow(ax, x + w, y + h / 2, nodes[i + 1][0], y + h / 2)

    _box(ax, 0.3, 0.45, 5.3, 1.2, "边端采图模块（既有）\n· 配置曝光/增益/行频/触发源=Line\n· 取流后写入本机图像文件", C_EDGE_FILL, C_EDGE, 8)
    _box(ax, 6.5, 0.45, 5.0, 1.2, "算法 V2（不控相机）\n· 不打开 SDK、不发软触发\n· 按本地文件路径读图（暂不共享内存）", C_ALGO_FILL, C_ALGO, 8)

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_triple_camera_flow(path: Path):
    """Three independent camera pipelines + board OR."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 5.8), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图4  三路相机独立检测与板尾汇总")

    cams = [
        (4.6, "上相机", "上棱胶缝 / 表面伤\n残胶 / 波浪 / 端头"),
        (2.8, "中相机", "侧面高低 / 开胶\n封边带损伤 / 崩缺"),
        (1.0, "下相机", "下棱胶缝 / 下表面伤\n残胶 / 端头"),
    ]
    for y, title, body in cams:
        _box(ax, 0.4, y, 1.8, 1.2, f"{title}\n线扫帧", C_HW_FILL, C_HW, 8.5)
        _arrow(ax, 2.2, y + 0.6, 2.7, y + 0.6)
        _box(ax, 2.7, y, 3.6, 1.2, f"独立流水线\n预处理→检测→分类→度量→判定\n{body}", C_ALGO_FILL, C_ALGO, 7.5)

    # merge
    _arrow(ax, 6.3, 5.2, 7.3, 3.5)
    _arrow(ax, 6.3, 3.4, 7.3, 3.3)
    _arrow(ax, 6.3, 1.6, 7.3, 3.1)
    _box(ax, 7.3, 2.5, 2.2, 1.6, "板尾汇总\nboard_ok = OR\n计数相加\n列表拼接", C_EDGE_FILL, C_EDGE, 9)
    _arrow(ax, 9.5, 3.3, 10.0, 3.3)
    _box(ax, 10.0, 2.6, 1.7, 1.4, "边端回传\nPLC / 中心", C_PLC_FILL, C_PLC, 8.5)

    ax.text(5.0, 0.35, "热路径不做跨相机同步融合 · 三路各跑各的 · 效率更高", ha="center", fontsize=8.5, color=C_ACCENT, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_single_pipeline(path: Path):
    """Single-channel detailed algorithm pipeline."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 5.2), dpi=160)
    ax.set_xlim(0, 12.5)
    ax.set_ylim(0, 5.5)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图5  单路算法处理流水线")

    # main chain
    boxes = [
        (0.2, "①读本地文件\n校验图像", C_ALGO_FILL),
        (2.2, "②预处理\nCLAHE/归一化", C_ALGO_FILL),
        (4.2, "③ROI / 门控\n空窗跳过", C_EDGE_FILL),
        (6.2, "④切片\n1024 窗口", C_EDGE_FILL),
        (8.2, "⑤检测\nYOLO11s@640", "#FFF3E0"),
        (10.2, "⑥分类\nYOLO11n-cls", "#FFF3E0"),
    ]
    y = 3.3
    for x, t, face in boxes:
        _box(ax, x, y, 1.85, 1.2, t, face, C_ALGO, 8)
    for i in range(len(boxes) - 1):
        x1 = boxes[i][0] + 1.85
        x2 = boxes[i + 1][0]
        _arrow(ax, x1, y + 0.6, x2, y + 0.6)

    # second row
    _arrow(ax, 11.1, 3.3, 9.6, 2.4)
    _arrow(ax, 9.6, 2.4, 7.0, 2.4)
    _box(ax, 5.2, 1.5, 2.2, 1.15, "⑦物理度量\n像素→mm", C_PLC_FILL, C_PLC, 8.5)
    _box(ax, 7.7, 1.5, 2.2, 1.15, "⑧阈值判定\n本路 OK/NG", C_EDGE_FILL, C_EDGE, 8.5)
    _box(ax, 10.2, 1.5, 2.0, 1.15, "输出缺陷列表\n+ 计数", C_ALGO_FILL, C_ALGO, 8.5)
    _arrow(ax, 7.4, 2.05, 7.7, 2.05)
    _arrow(ax, 9.9, 2.05, 10.2, 2.05)

    # models note
    _box(
        ax,
        0.3,
        0.25,
        4.5,
        1.0,
        "模型组合：检测 YOLO11s（~9.4M）+ 分类 YOLO11n\n加速：TensorRT FP16 · 配方 4 套 det+cls 热加载",
        C_BG,
        "#888888",
        7.5,
    )
    _box(
        ax,
        5.1,
        0.25,
        7.0,
        1.0,
        "性能目标（单路，不含采图）：预处理+门控 ≤5ms · det ≤12ms · cls ≤4ms · 度量判定 ≤2ms · 合计 ≤23ms（原 2080 Ti 等配置裕度更宽）",
        C_BG,
        "#888888",
        7.5,
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_sample_workflow(path: Path):
    """Sample collection and labeling loop."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4.8)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图6  样本收集 · 标注 · 训练闭环")

    items = [
        (0.3, "产线采图\n三路·多配方", C_HW_FILL, C_HW),
        (2.6, "缺陷真值确认\n工艺/质检", C_PLC_FILL, C_PLC),
        (4.9, "标注 bbox\n检测+分类集", C_EDGE_FILL, C_EDGE),
        (7.2, "训练 det/cls\nTensorRT 部署", C_ALGO_FILL, C_ALGO),
        (9.5, "试运行回流\n漏检/误检补标", C_ACCENT, "#F5D5C0"),
    ]
    for i, (x, t, face, edge) in enumerate(items):
        _box(ax, x, 2.0, 2.1, 1.4, t, face if i < 4 else "#F5D5C0", edge if i < 4 else C_ACCENT, 8.5)
        if i < len(items) - 1:
            _arrow(ax, x + 2.1, 2.7, items[i + 1][0], 2.7)

    # feedback loop arrow
    ax.annotate(
        "",
        xy=(1.35, 2.0),
        xytext=(10.55, 2.0),
        arrowprops=dict(arrowstyle="-|>", color=C_ACCENT, lw=1.5, connectionstyle="arc3,rad=-0.35"),
    )
    ax.text(6.0, 0.7, "难例回流闭环：持续提升检出率与稳定性（必要支撑，非可选）", ha="center", fontsize=9, color=C_ACCENT, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        from docx.shared import RGBColor

        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, name="黑体", size=16 if level == 1 else 14 if level == 2 else 12, bold=True)
    return h


def add_para(doc, text, size=11, bold=False, first_line_indent=True):
    p = doc.add_paragraph()
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=9, bold=True)
    return p


def add_figure(doc, image_path: Path, width_inches=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, name="黑体", size=10, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def generate_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    paths = {
        "arch": FIG_DIR / "fig1_system_architecture.png",
        "boundary": FIG_DIR / "fig2_module_boundary.png",
        "acq": FIG_DIR / "fig3_acquisition_flow.png",
        "triple": FIG_DIR / "fig4_triple_camera.png",
        "pipe": FIG_DIR / "fig5_single_pipeline.png",
        "sample": FIG_DIR / "fig6_sample_workflow.png",
    }
    draw_system_architecture(paths["arch"])
    draw_module_boundary(paths["boundary"])
    draw_acquisition_flow(paths["acq"])
    draw_triple_camera_flow(paths["triple"])
    draw_single_pipeline(paths["pipe"])
    draw_sample_workflow(paths["sample"])
    return paths


def main():
    figs = generate_figures()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    run = title.add_run("板材封边缺陷智能检测设备")
    set_run_font(run, name="黑体", size=22, bold=True)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = st.add_run("视觉算法架构说明书（V2）")
    set_run_font(run, name="黑体", size=20, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(20)
    run = meta.add_run(
        "文档版本：V2.3\n用途：算法架构说明、与现有软件模块边界、样本支撑要求\n"
        "含系统架构图、模块边界图；图像落盘传图；算法OS优先Ubuntu"
    )
    set_run_font(run, size=11)

    doc.add_page_break()

    # 1
    add_heading_cn(doc, "1 概述", 1)
    add_heading_cn(doc, "1.1 文档目的", 2)
    add_para(
        doc,
        "本说明书描述板材封边缺陷视觉算法 V2 的架构设计结果，明确算法服务与现有边端客户端、采图模块、"
        "视觉中心、PLC、MES 等软件模块的职责边界，说明算法流水线与三路相机处理策略，并给出样本收集与标注支撑要求。",
    )
    add_heading_cn(doc, "1.2 设计定位", 2)
    add_para(
        doc,
        "在现有边端客户端 / 视觉中心 / PLC / MES 体系上升级边端 AI 检测算法；不改造 MES、贴标与出料中控协议。"
        "算法以边端同机推理服务部署；边端提交图像本机文件路径与配方，算法读盘后回传结果"
        "（现阶段不使用共享内存）。算法 OS 优先 Ubuntu，迁移动成本高时可与边端同沿用原系统。"
        "三路线扫各跑各的，热路径不做跨相机同步融合。",
    )
    add_heading_cn(doc, "1.3 设计目标", 2)
    add_table(
        doc,
        ["项", "设计结果"],
        [
            ["线速", "最高 60 m/min"],
            ["像素精度", "约 0.03 mm 量级（以现场标定为准）"],
            ["检出目标", "板级缺陷检出率 ≥99%（验收口径以合同为准）"],
            ["软件边界", "保留现有边端 UI、视觉中心、PLC、MES；仅替换边端「AI 推理」步骤"],
            ["硬件策略", "先沿用现场原硬件配置复原功能；功能稳定后再评估优化降成本（如 GPU）"],
            ["异常策略", "算法服务不可用时，边端回退原模型槽位，产线不停线"],
        ],
    )

    # 2 Architecture rich
    add_heading_cn(doc, "2 总体架构", 1)
    add_heading_cn(doc, "2.1 架构说明", 2)
    add_para(
        doc,
        "系统分为电气触发层、成像层、边端既有软件层、算法 V2 层与产线上层五大部分。"
        "电气侧完成硬触发，相机完成成像，边端采图模块完成 SDK 参数与取流，算法 V2 完成检测与判定，"
        "结果仍由边端按原协议下发 PLC、上报视觉中心与 MES。图1 给出端到端全景关系。",
    )
    add_figure(doc, figs["arch"], 6.4)
    add_caption(doc, "图1  系统总体架构（硬件 / 边端既有软件 / 算法 V2 / 上层系统）")

    add_heading_cn(doc, "2.2 分层职责", 2)
    add_table(
        doc,
        ["层级", "模块", "职责"],
        [
            ["电气", "光电、编码器、PLC", "产生硬触发，决定曝光时刻"],
            ["成像", "三路线扫相机", "按 TRIG 曝光，输出 Mono8"],
            ["边端现有", "采图模块 + UI + 通讯", "相机参数与取流；业务交互；结果外发"],
            ["边端新增", "视觉算法服务 V2", "缺陷检测 / 分类 / 度量 / 判定"],
            ["产线上层", "视觉中心、MES", "多机汇总、复检、业务数据（协议不变）"],
        ],
    )
    add_para(
        doc,
        "部署要点：边端采图落盘 → 下发本机路径 → V2 读文件推理；控制与元数据走本地接口；"
        "算法优先 Ubuntu（可沿用 Windows）；先原硬件复原，再评估降本；视觉中心不跑 V2。",
        first_line_indent=True,
    )

    # 3 boundaries rich
    add_heading_cn(doc, "3 算法 V2 与其他软件模块边界", 1)
    add_para(
        doc,
        "分界原则：能由现场既有软件与电气完成的环节不纳入 V2；V2 按本机文件路径读取已落盘图像并推理，"
        "将结构化结果交还边端。图2 以泳道方式标明各方职责与数据方向。",
    )
    add_figure(doc, figs["boundary"], 6.4)
    add_caption(doc, "图2  算法 V2 与周边软件模块边界")

    add_heading_cn(doc, "3.1 边界总表", 2)
    add_table(
        doc,
        ["模块", "归属", "负责内容", "不负责内容"],
        [
            [
                "视觉算法服务 V2",
                "本方案新增",
                "按文件路径读图；预处理/切片；检测与分类；物理度量；单路判定；板尾 OR 汇总",
                "相机 SDK/曝光/硬触发；PLC 点位；MES/贴标协议；边端 UI/加密狗",
            ],
            [
                "边端客户端",
                "既有",
                "工程方案/ROI/阈值；图像落盘；下发路径调用 V2；展示与转发；失败回退原模型",
                "V2 内部推理实现",
            ],
            [
                "边端采图模块",
                "边端内",
                "SDK 曝光与硬触发取流；组帧；写本机图像文件",
                "缺陷识别与业务判定",
            ],
            [
                "PLC/光电/编码器",
                "电气控制",
                "硬触发；运动相关信号；接收 NG 等控制",
                "图像算法",
            ],
            [
                "视觉中心",
                "既有",
                "多机汇总、复检、贴标路径协同",
                "边端本地 V2 推理",
            ],
            [
                "MES",
                "既有",
                "业务数据交互（原协议）",
                "模型与采图控制",
            ],
        ],
    )

    add_heading_cn(doc, "3.2 采图与硬触发边界", 2)
    add_para(
        doc,
        "线扫需上位机配置曝光并与硬触发配合取流。硬触发图像获取由「边端采图模块 + PLC」完成；"
        "边端将图像落盘后，V2 **按本地文件路径**读图（现阶段不使用共享内存，见图3）。",
    )
    add_figure(doc, figs["acq"], 6.4)
    add_caption(doc, "图3  采图与硬触发协作流程")
    add_table(
        doc,
        ["环节", "负责方", "设计结果"],
        [
            ["曝光/增益/行频/触发模式", "边端采图模块（SDK）", "工程方案配置并下发相机"],
            ["硬触发脉冲", "PLC → 相机 TRIG", "按电气时序曝光；禁用软触发量产"],
            ["取流与落盘", "边端采图/客户端", "写本机约定目录图像文件"],
            ["算法输入", "边端 → V2", "本机文件绝对路径 + 板材/配方信息"],
            ["帧接入", "算法 V2", "按路径读文件；不打开相机、不发 TRIG"],
        ],
    )

    add_heading_cn(doc, "3.3 接口边界摘要", 2)
    add_para(
        doc,
        "边端 → V2：板材 ID、机台号、配方 ID、**三路图像本机路径**、ROI、标定（mm/px）、判定阈值。",
        first_line_indent=False,
    )
    add_table(
        doc,
        ["字段（V2 → 边端）", "说明"],
        [
            ["board_ok", "板级 OK / NG"],
            ["defects[]", "类别、相机位、框、宽/长/面积（mm）、置信度"],
            ["counts", "各类计数"],
            ["engine", "标识 scanly-v2"],
        ],
    )
    add_para(doc, "边端再按原协议写入 PLC、上报视觉中心。V2 不直连 PLC / MES。")

    # 4 triple camera
    add_heading_cn(doc, "4 三路相机策略", 1)
    add_para(
        doc,
        "上、中、下三路主责不同成像区域，业务上无需跨相机框级融合。每路独立完成检测与判定，"
        "板材离开后仅做 OR / 计数汇总（见图4），利于 GPU 利用率并避免跨路等待。",
    )
    add_figure(doc, figs["triple"], 6.4)
    add_caption(doc, "图4  三路相机独立检测与板尾汇总")
    add_table(
        doc,
        ["项", "设计结果"],
        [
            ["运行方式", "上 / 中 / 下三路独立检测、独立判定"],
            ["板级结论", "任一路 NG → 板 NG；列表拼接；计数相加"],
            ["跨相机融合", "不做框级同步融合与跨路等待"],
        ],
    )
    add_table(
        doc,
        ["相机", "主检内容"],
        [
            ["上", "上棱胶缝、上表面划伤/压痕、残胶、波浪纹；端头过长/短带（上缘可见）"],
            ["中", "侧面低于/高于饰面、开胶、封边带损伤、崩缺"],
            ["下", "下棱胶缝、下表面伤、残胶；端头过长/短带（下缘可见）"],
        ],
    )

    # 5 pipeline rich
    add_heading_cn(doc, "5 单路算法流程与模型", 1)
    add_para(
        doc,
        "单路在边端交帧后按固定流水线处理。检测模型保证定位召回，分类模型对裁剪区域做细分类与误报抑制，"
        "再结合物理标定尺寸与可配置阈值给出本路 OK/NG（见图5）。",
    )
    add_figure(doc, figs["pipe"], 6.5)
    add_caption(doc, "图5  单路算法处理流水线")

    add_heading_cn(doc, "5.1 步骤说明", 2)
    add_table(
        doc,
        ["步骤", "内容", "说明"],
        [
            ["① 读本地文件", "按路径打开已落盘图像", "不控相机；暂不共享内存"],
            ["② 预处理", "CLAHE / 归一化等 GPU 处理", "减轻后续网络负担"],
            ["③ ROI/门控", "封边区域 ROI；空窗跳过", "降低无效推理"],
            ["④ 切片", "1024 窗口，重叠可配", "应对极小目标"],
            ["⑤ 检测", "YOLO11s，输入默认 640", "高召回定位"],
            ["⑥ 分类", "YOLO11n-cls，裁剪约 224", "细分类与打假阳"],
            ["⑦ 度量", "像素→毫米", "使用标定 mm/px"],
            ["⑧ 判定", "可配置阈值", "输出本路结果"],
        ],
    )
    add_heading_cn(doc, "5.2 模型配置", 2)
    add_table(
        doc,
        ["模型", "规格", "作用"],
        [
            ["检测", "YOLO11s，输入默认 640", "定位缺陷区域"],
            ["分类", "YOLO11n-cls，裁剪约 224", "细分类与误报抑制"],
            ["加速", "TensorRT FP16", "边端实时推理"],
            ["配方", "4 套（浅/深板 × 浅/深封边）各一组 det+cls", "与工程方案对应热加载"],
        ],
    )

    # 6 thresholds
    add_heading_cn(doc, "6 缺陷判定（默认阈值，可配置）", 1)
    add_table(
        doc,
        ["缺陷", "默认条件（mm）"],
        [
            ["胶缝", "宽 >0.2 且 长 >5"],
            ["侧面低于饰面", "高 >0.3 且 长 >5"],
            ["短带", "长 >1"],
            ["过长", "长 >1"],
            ["划伤/压痕", "宽 >0.4 且 长 >5"],
            ["残胶", "宽 >1 且 长 >3"],
            ["开胶", "宽 >1 且 长 >5（可配置为检出即 NG）"],
            ["崩缺", "宽 >1 且 长 >1"],
            ["封边带损伤", "宽 >1 且 长 >1"],
            ["波浪纹", "宽 >0.7"],
        ],
    )
    add_para(doc, "阈值由工程方案下发，支持现场调整，不写死在模型内。")

    # 7 samples
    add_heading_cn(doc, "7 样本收集与标注支撑要求", 1)
    add_para(
        doc,
        "视觉算法 V2 的检出率与稳定性高度依赖现场真实样本。模型训练、验证、配方适配及持续迭代均需双方协同"
        "完成样本收集与标注。本项为算法交付与验收的必要支撑条件。图6 给出数据闭环。",
    )
    add_figure(doc, figs["sample"], 6.4)
    add_caption(doc, "图6  样本收集 · 标注 · 训练闭环")

    add_heading_cn(doc, "7.1 样本范围与规模", 2)
    add_table(
        doc,
        ["类别", "要求"],
        [
            ["缺陷类型", "覆盖约定缺陷：胶缝、侧面高低、短带/过长、划伤、残胶、开胶、崩缺、封边带损伤、波浪纹等"],
            ["色系配方", "浅/深板 × 浅/深封边等实际配方均需代表性样本"],
            ["相机视角", "上 / 中 / 下三路均需采集并可区分"],
            ["规模建议", "每类有效缺陷框建议 ≥1000 量级（极明显类可酌减）"],
            ["负样本/难例", "正常板面与易误报场景；试运行漏检误检持续回流"],
        ],
    )
    add_heading_cn(doc, "7.2 标注与协作分工", 2)
    add_para(
        doc,
        "标注以水平矩形框（bbox）为主；检测集与分类裁剪集分别建设。标注规范与争议样例需双方确认业务口径。",
    )
    add_table(
        doc,
        ["事项", "甲方（客户/现场）", "乙方（算法方）"],
        [
            ["产线采图", "提供采图条件或授权存图，覆盖配方与班次", "提出采集清单与质量要求"],
            ["缺陷真值", "工艺/质检确认；协助难例收集", "类别体系与标注规范"],
            ["标注执行", "安排或委托标注；参与抽检与口径确认", "工具/规范、质检与清洗"],
            ["迭代", "持续提供漏检误检回流", "补训与版本发布"],
            ["验收数据", "确认验收集代表性", "在约定数据集报告指标"],
        ],
    )

    # 8 hardware
    add_heading_cn(doc, "8 硬件配置与性能", 1)
    add_heading_cn(doc, "8.1 实施策略", 2)
    add_para(
        doc,
        "算法 V2 的硬件原则：优先沿用厦门四期 / 现场既有硬件配置，先完成功能复原与联调验收；"
        "功能与检出指标稳定后，再单独评估硬件优化与降本（例如评估 RTX 3060 等更经济 GPU 是否满足节拍）。"
        "不以更换相机或升级 GPU 作为功能上线前置条件。",
    )
    add_heading_cn(doc, "8.2 沿用原硬件配置（建议基线）", 2)
    add_para(
        doc,
        "以下配置依据既有软件设计说明书与物料清单归纳，作为现阶段复原与联调的推荐基线"
        "（具体以现场实机型号为准）。",
    )
    add_table(
        doc,
        ["类别", "配置（原方案/现场常用）", "说明"],
        [
            ["线扫相机", "3× 海康 MV-CL024-91GM（上/中/下）", "GigE · 硬触发 · Mono8"],
            ["镜头/光源", "工业镜头 + 条形/线性光源及控制器", "按既有机位保留，算法不改光学"],
            ["触发链路", "编码器 + 光电 → PLC → 相机 TRIG", "量产硬触发，禁用软触发"],
            ["边端工控机", "i7 级 CPU，16GB 起内存，SSD 2TB 级", "如原 HC-SYSA35II 等机型"],
            ["边端 GPU", "RTX 2080 Ti（现场/原方案常用）", "现阶段主用，保证复原算力裕度"],
            ["算法 OS", "优先 Ubuntu Linux", "切 Linux 成本过高时可与边端同沿用 Windows"],
            ["图像传递", "本机落盘 + 文件绝对路径", "现阶段不使用共享内存，利于解耦与复盘"],
            ["视觉中心", "汇总/复检主机", "不运行 V2 推理；内存可比边端更高"],
            ["网络", "工业千兆网卡 / 交换机", "三相机 + 边端稳定取流"],
            ["通讯", "PLC Modbus TCP / S7；MES HTTP", "协议保持不变"],
        ],
    )
    add_heading_cn(doc, "8.3 后续优化方向（非当前前置）", 2)
    add_table(
        doc,
        ["方向", "内容", "前提"],
        [
            ["GPU 降本评估", "在模型与门控优化后，评估 RTX 3060 12GB 等是否满足节拍", "功能已在原硬件上稳定达标"],
            ["工控机精简", "在算力证明过剩时再下调 CPU/内存规格", "连续跑线与散热验证完成"],
            ["光学微调", "仅在成像质量不足时再调整光机", "先排除采图参数与触发问题"],
        ],
    )
    add_heading_cn(doc, "8.4 性能目标（算法侧）", 2)
    add_table(
        doc,
        ["项", "目标"],
        [
            ["单路处理（不含采图/盘 IO 视磁盘而定）", "算法段 ≤23 ms 量级目标（落盘时间由边端与磁盘共同承担）"],
            ["线速目标", "最高 60 m/min（以产线与验收约定为准，可分档验证）"],
            ["视觉中心", "不运行本算法推理"],
        ],
    )

    # 9 phases
    add_heading_cn(doc, "9 实施阶段", 1)
    add_table(
        doc,
        ["阶段", "交付内容"],
        [
            [
                "一阶段（原硬件复原）",
                "在既有相机/工控机/2080 Ti 等原配置上：单路检测+分类+度量；与边端采图/回退对接；基础样本训练集就绪",
            ],
            [
                "二阶段（三路与配方）",
                "三路独立并行；板尾 OR 汇总；四配方切换；与视觉中心对齐；配方样本补齐",
            ],
            [
                "三阶段（达标与优化评估）",
                "参数固化与目标线速连续运行验收；难例回流补训；可选评估硬件降本，不倒逼更换设备",
            ],
        ],
    )

    add_heading_cn(doc, "10 附则", 1)
    add_para(
        doc,
        "本说明书描述已确认的设计结果。接口字段细节与内部实现以双方联调约定及后续技术附件为准。"
        "图件与正文一致；若硬件、软件版本或缺陷口径变更，应同步评估对架构边界与样本计划的影响。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote figures under: {FIG_DIR}")
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    main()
