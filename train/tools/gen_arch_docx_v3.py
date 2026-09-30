#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Generate Scanly Vision Algo V3 architecture Word specification with diagrams.

V3 relative to V2:
1. Area-scan (2×) cold path for material/pattern/tone → recipe / exposure / thresholds.
2. Inference runtime fixed to Windows IoT (not Ubuntu-preferred).
Other content reuses V2 design conclusion.
"""

from __future__ import annotations

from pathlib import Path

import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt

ROOT = Path(__file__).resolve().parents[2]
DOC_DIR = ROOT / "doc" / "02-dr" / "视觉算法"
FIG_DIR = DOC_DIR / "figures_v3"
OUT = DOC_DIR / "板材封边缺陷视觉算法架构说明书V3.docx"

# industrial palette
C_BG = "#F5F7FA"
C_EDGE = "#1F4E79"
C_EDGE_FILL = "#D6E3F0"
C_ALGO = "#0D7377"
C_ALGO_FILL = "#D5F0F1"
C_HW = "#5B5B5B"
C_HW_FILL = "#E8E8E8"
C_PLC = "#8B5A00"
C_PLC_FILL = "#F5E6C8"
C_ARROW = "#333333"
C_TEXT = "#1A1A1A"
C_ACCENT = "#C45C26"


def _setup_cn_font():
    plt.rcParams["font.sans-serif"] = [
        "PingFang SC",
        "Heiti SC",
        "STHeiti",
        "Songti SC",
        "Arial Unicode MS",
        "DejaVu Sans",
    ]
    plt.rcParams["axes.unicode_minus"] = False


def _box(ax, x, y, w, h, text, face, edge, fontsize=9, lw=1.4):
    patch = FancyBboxPatch(
        (x, y),
        w,
        h,
        boxstyle="round,pad=0.02,rounding_size=0.08",
        facecolor=face,
        edgecolor=edge,
        linewidth=lw,
        mutation_aspect=0.3,
    )
    ax.add_patch(patch)
    ax.text(
        x + w / 2,
        y + h / 2,
        text,
        ha="center",
        va="center",
        fontsize=fontsize,
        color=C_TEXT,
        linespacing=1.25,
        wrap=True,
    )
    return patch


def _arrow(ax, x1, y1, x2, y2, text="", color=C_ARROW):
    ax.annotate(
        "",
        xy=(x2, y2),
        xytext=(x1, y1),
        arrowprops=dict(arrowstyle="-|>", color=color, lw=1.3, mutation_scale=12),
    )
    if text:
        mx, my = (x1 + x2) / 2, (y1 + y2) / 2
        ax.text(mx, my + 0.08, text, ha="center", va="bottom", fontsize=7.5, color=color)


def _subtitle(ax, text, x=0.5, y=0.97):
    ax.text(x, y, text, transform=ax.transAxes, ha="center", va="top", fontsize=11, fontweight="bold", color=C_EDGE)


def draw_system_architecture(path: Path):
    """Overall system V3: area+line cameras → recipe → algo V3; Windows IoT."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 7.8), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8.5)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    ax.set_facecolor("white")
    _subtitle(ax, "图1  系统总体架构（V3）")

    # left hardware
    _box(ax, 0.25, 7.2, 2.3, 0.75, "光电 / 编码器", C_HW_FILL, C_HW, 8.5)
    _box(ax, 0.25, 6.15, 2.3, 0.75, "PLC 触发\n面扫进板 / 线扫 TRIG", C_PLC_FILL, C_PLC, 8)
    _box(ax, 0.25, 4.85, 2.3, 1.0, "2× 面扫相机\n材质/花纹/色系", C_HW_FILL, C_HW, 8)
    _box(ax, 0.25, 3.35, 2.3, 1.2, "3× 线扫相机\n上 / 中 / 下\nGigE · Mono8", C_HW_FILL, C_HW, 8)
    _arrow(ax, 1.4, 7.2, 1.4, 6.9)
    _arrow(ax, 1.4, 6.15, 1.4, 5.85)
    _arrow(ax, 1.4, 4.85, 1.4, 4.55)

    # edge existing group
    group = FancyBboxPatch(
        (2.9, 2.55),
        4.0,
        5.4,
        boxstyle="round,pad=0.03,rounding_size=0.1",
        facecolor="#F0F4F8",
        edgecolor=C_EDGE,
        linewidth=1.6,
    )
    ax.add_patch(group)
    ax.text(4.9, 7.7, "边端客户端（现有，保留）", ha="center", va="center", fontsize=9.5, fontweight="bold", color=C_EDGE)
    _box(ax, 3.15, 6.55, 3.5, 0.85, "工程方案 / ROI / 阈值\n配方管理（可扩展 N 套）", C_EDGE_FILL, C_EDGE, 8)
    _box(ax, 3.15, 5.35, 3.5, 0.95, "采图模块\n面扫+线扫取流 · 落盘\n按配方写曝光/增益", C_EDGE_FILL, C_EDGE, 7.5)
    _box(ax, 3.15, 4.15, 3.5, 0.95, "配方决议\n面扫识别 + MES 冲突策略", "#E8F0E0", C_ALGO, 8)
    _box(ax, 3.15, 2.85, 3.5, 1.05, "通讯与业务\n→ PLC · 视觉中心 / MES", C_EDGE_FILL, C_EDGE, 8)
    _arrow(ax, 4.9, 6.55, 4.9, 6.3)
    _arrow(ax, 4.9, 5.35, 4.9, 5.1)
    _arrow(ax, 4.9, 4.15, 4.9, 3.9)

    _arrow(ax, 2.55, 5.35, 3.15, 5.9, "面扫图")
    _arrow(ax, 2.55, 3.95, 3.15, 5.55, "线扫图")

    # algo V3
    group2 = FancyBboxPatch(
        (7.4, 2.55),
        4.3,
        5.4,
        boxstyle="round,pad=0.03,rounding_size=0.1",
        facecolor="#EAF7F7",
        edgecolor=C_ALGO,
        linewidth=1.6,
    )
    ax.add_patch(group2)
    ax.text(9.55, 7.7, "视觉算法服务 V3（本机 · Windows IoT）", ha="center", va="center", fontsize=9, fontweight="bold", color=C_ALGO)

    steps = [
        (6.75, "面扫色系/花纹识别（冷路径）"),
        (5.95, "帧接入_线扫本地文件"),
        (5.15, "预处理 / ROI / 切片"),
        (4.35, "检测 YOLO11s"),
        (3.55, "分类 YOLO11n-cls"),
        (2.75, "度量+判定 → 板尾 OR 汇总"),
    ]
    for y, t in steps:
        face = "#FFF3E0" if y == 6.75 else C_ALGO_FILL
        _box(ax, 7.65, y, 3.85, 0.62, t, face, C_ALGO, 8)
    for i in range(len(steps) - 1):
        y1 = steps[i][0]
        y2 = steps[i + 1][0] + 0.62
        _arrow(ax, 9.55, y1, 9.55, y2)

    _arrow(ax, 6.9, 4.85, 7.65, 6.1, "面扫路径")
    _arrow(ax, 6.9, 5.7, 7.65, 5.4, "线扫路径")
    _arrow(ax, 6.9, 6.9, 7.65, 6.95, "recipe/阈值")
    _arrow(ax, 7.65, 2.95, 6.9, 3.2, "OK/NG")

    ax.text(
        6.0,
        0.55,
        "部署：边端同机 Windows IoT · 落盘+本地路径 · 2面扫配方冷路径 + 3线扫缺陷热路径 · 配方 N 套可扩展 · 不改中心/MES 协议",
        ha="center",
        va="center",
        fontsize=7.2,
        color="#555555",
        style="italic",
    )
    ax.text(
        6.0,
        0.15,
        "相对 V2：合入面扫设计；推理环境固定 Windows IoT（非 Ubuntu 优先）",
        ha="center",
        va="center",
        fontsize=7.5,
        color=C_ACCENT,
        fontweight="bold",
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_module_boundary(path: Path):
    """Responsibility swimlane-style boundary diagram."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 7.2), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 8.0)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图2  算法 V3 与周边软件模块边界")

    # zone banners (no overlap with body)
    zones = [
        (0.2, 4.4, "现场既有 · 电气", C_PLC_FILL, C_PLC),
        (2.5, 6.7, "边端软件（采图 + 客户端 + 面扫识别 + 算法 V3）", C_EDGE_FILL, C_EDGE),
        (9.4, 2.4, "产线上层", C_HW_FILL, C_HW),
    ]
    for x, w, title, face, edge in zones:
        _box(ax, x, 6.95, w, 0.45, title, face, edge, 8, lw=1.2)

    lanes = [
        (0.2, "电气 / PLC", C_PLC_FILL, C_PLC, "硬触发脉冲\n进板/到位信号\n接收 NG 控制"),
        (2.5, "边端采图", C_EDGE_FILL, C_EDGE, "面扫+线扫 SDK\n曝光下发\n图像落盘"),
        (4.8, "边端客户端\n+配方决议", C_EDGE_FILL, C_EDGE, "工程方案/ROI\n面扫→recipe\n路径与阈值下发"),
        (7.1, "算法 V3", C_ALGO_FILL, C_ALGO, "面扫分类冷路径\n线扫 det/cls\nWindows IoT 部署"),
        (9.4, "中心 / MES", C_HW_FILL, C_HW, "多机汇总复检\n色系先验可校\n协议不变"),
    ]
    for x, title, face, edge, body in lanes:
        _box(ax, x, 5.85, 2.1, 0.7, title, face, edge, 8.5, lw=1.6)
        _box(ax, x, 3.0, 2.1, 2.55, body, "#FFFFFF", edge, 8.5)

    # flow arrows between body columns
    for x in [2.3, 4.6, 6.9, 9.2]:
        _arrow(ax, x, 4.25, x + 0.2, 4.25)

    ax.text(6.0, 2.25, "数据方向（示意）", ha="center", fontsize=9, fontweight="bold", color=C_TEXT)
    flow_items = [
        "TRIG / 进板",
        "落盘图像",
        "路径+recipe",
        "OK/NG+缺陷",
        "原协议上报",
    ]
    xs = [1.25, 3.55, 5.85, 8.15, 10.45]
    for x, t in zip(xs, flow_items):
        _box(ax, x - 0.9, 1.05, 1.8, 0.75, t, C_BG, "#888888", 7.5)
    for i in range(len(xs) - 1):
        _arrow(ax, xs[i] + 0.9, 1.4, xs[i + 1] - 0.9, 1.4)

    _box(
        ax,
        0.3,
        0.15,
        11.4,
        0.65,
        "说明：V3 合入 2 面扫色系识别冷路径；线扫缺陷热路径同 V2；算法服务部署于 Windows IoT（边端同机）",
        "#FFF8F0",
        C_ACCENT,
        8,
        lw=1.0,
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_acquisition_flow(path: Path):
    """Hard trigger acquisition sequence."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 4.8), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图3  采图与硬触发协作流程")

    nodes = [
        (0.3, "光电检测到板"),
        (2.5, "PLC 输出 TRIG"),
        (4.7, "相机按设定曝光"),
        (6.9, "边端 SDK 取流并落盘"),
        (9.1, "V3 按路径读图\n帧接入"),
    ]
    y = 2.4
    w, h = 1.9, 1.15
    for i, (x, t) in enumerate(nodes):
        face = C_ALGO_FILL if i == 4 else (C_PLC_FILL if i == 1 else C_EDGE_FILL if i >= 3 else C_HW_FILL)
        edge = C_ALGO if i == 4 else (C_PLC if i == 1 else C_EDGE if i >= 3 else C_HW)
        _box(ax, x, y, w, h, t, face, edge, 8.5)
        if i < len(nodes) - 1:
            _arrow(ax, x + w, y + h / 2, nodes[i + 1][0], y + h / 2)

    _box(ax, 0.3, 0.45, 5.3, 1.2, "边端采图模块（既有）\n· 配置曝光/增益/行频/触发源=Line\n· 取流后写入本机图像文件", C_EDGE_FILL, C_EDGE, 8)
    _box(ax, 6.5, 0.45, 5.0, 1.2, "算法 V3（不控相机）\n· 不打开 SDK、不发软触发\n· Windows IoT 上按路径读图（暂不共享内存）", C_ALGO_FILL, C_ALGO, 8)

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_triple_camera_flow(path: Path):
    """Three independent camera pipelines + board OR."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 5.8), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 6.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图4  三路相机独立检测与板尾汇总")

    cams = [
        (4.6, "上相机", "上棱胶缝 / 表面伤\n残胶 / 波浪 / 端头"),
        (2.8, "中相机", "侧面高低 / 开胶\n封边带损伤 / 崩缺"),
        (1.0, "下相机", "下棱胶缝 / 下表面伤\n残胶 / 端头"),
    ]
    for y, title, body in cams:
        _box(ax, 0.4, y, 1.8, 1.2, f"{title}\n线扫帧", C_HW_FILL, C_HW, 8.5)
        _arrow(ax, 2.2, y + 0.6, 2.7, y + 0.6)
        _box(ax, 2.7, y, 3.6, 1.2, f"独立流水线\n预处理→检测→分类→度量→判定\n{body}", C_ALGO_FILL, C_ALGO, 7.5)

    # merge
    _arrow(ax, 6.3, 5.2, 7.3, 3.5)
    _arrow(ax, 6.3, 3.4, 7.3, 3.3)
    _arrow(ax, 6.3, 1.6, 7.3, 3.1)
    _box(ax, 7.3, 2.5, 2.2, 1.6, "板尾汇总\nboard_ok = OR\n计数相加\n列表拼接", C_EDGE_FILL, C_EDGE, 9)
    _arrow(ax, 9.5, 3.3, 10.0, 3.3)
    _box(ax, 10.0, 2.6, 1.7, 1.4, "边端回传\nPLC / 中心", C_PLC_FILL, C_PLC, 8.5)

    ax.text(5.0, 0.35, "热路径不做跨相机同步融合 · 三路各跑各的 · 效率更高", ha="center", fontsize=8.5, color=C_ACCENT, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def draw_single_pipeline(path: Path):
    """Single-channel pipeline: ⑥ 正下方竖直接入 ⑦（无斜线/折线）。"""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 5.4), dpi=160)
    ax.set_xlim(0, 12.5)
    ax.set_ylim(0, 5.5)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图5  单路算法处理流水线")

    bw, bh = 1.42, 1.1
    y_top = 3.48
    b2w, b2h, gap2 = 1.28, 1.05, 0.16
    y_bot = 1.55
    gap1 = 0.14

    # 上行 6 步自左向右；⑥ 中心即竖落点
    x0 = 0.35
    xs = [x0 + i * (bw + gap1) for i in range(6)]
    drop_x = xs[-1] + bw / 2

    labels = [
        "①读本地文件\n校验图像",
        "②预处理\nCLAHE/归一化",
        "③ROI / 门控\n空窗跳过",
        "④切片\n1024 窗口",
        "⑤检测\nYOLO11s@640",
        "⑥分类\nYOLO11n-cls",
    ]
    faces = [C_ALGO_FILL, C_ALGO_FILL, C_EDGE_FILL, C_EDGE_FILL, "#FFF3E0", "#FFF3E0"]
    for x, t, face in zip(xs, labels, faces):
        _box(ax, x, y_top, bw, bh, t, face, C_ALGO, 7.5)
    for i in range(5):
        _arrow(ax, xs[i] + bw, y_top + bh / 2, xs[i + 1], y_top + bh / 2)

    # 下行：⑦ 中心对齐 drop_x；⑧、输出向右
    x7 = drop_x - b2w / 2
    x8 = x7 + b2w + gap2
    x9 = x8 + b2w + gap2
    if x9 + b2w > 11.95:
        # 右侧越界时略缩 gap，保证输出框完整；⑦ 仍对准 ⑥
        gap2 = max(0.1, gap2 - (x9 + b2w - 11.95) / 2)
        x8 = x7 + b2w + gap2
        x9 = x8 + b2w + gap2

    _box(ax, x7, y_bot, b2w, b2h, "⑦物理度量\n像素→mm", C_PLC_FILL, C_PLC, 8)
    _box(ax, x8, y_bot, b2w, b2h, "⑧阈值判定\n本路 OK/NG", C_EDGE_FILL, C_EDGE, 8)
    _box(ax, x9, y_bot, b2w, b2h, "输出缺陷列表\n+ 计数", C_ALGO_FILL, C_ALGO, 8)
    _arrow(ax, x7 + b2w, y_bot + b2h / 2, x8, y_bot + b2h / 2)
    _arrow(ax, x8 + b2w, y_bot + b2h / 2, x9, y_bot + b2h / 2)

    # ⑥ 底 → ⑦ 顶：纯竖直短线 + 箭头（不用斜连/折线）
    y_tip = y_bot + b2h
    ax.plot([drop_x, drop_x], [y_top, y_tip + 0.08], color=C_ARROW, lw=1.45, solid_capstyle="butt", zorder=2)
    ax.annotate(
        "",
        xy=(drop_x, y_tip + 0.01),
        xytext=(drop_x, y_tip + 0.2),
        arrowprops=dict(arrowstyle="-|>", color=C_ARROW, lw=0, mutation_scale=12),
        zorder=3,
    )

    _box(
        ax,
        0.25,
        0.2,
        4.7,
        1.0,
        "模型组合：检测 YOLO11s（~9.4M）+ 分类 YOLO11n\n加速：TensorRT FP16 · 配方 N 套配置化热加载（不锁 4）",
        C_BG,
        "#888888",
        7.5,
    )
    _box(
        ax,
        5.15,
        0.2,
        6.8,
        1.0,
        "性能目标（单路，不含采图）：预处理+门控 ≤5ms · det ≤12ms · cls ≤4ms · 度量判定 ≤2ms · 合计 ≤23ms（原 2080 Ti 等配置裕度更宽）",
        C_BG,
        "#888888",
        7.5,
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)



def draw_sample_workflow(path: Path):
    """Sample collection and labeling loop."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.2, 4.6), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 4.8)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图6  样本收集 · 标注 · 训练闭环")

    items = [
        (0.3, "产线采图\n面扫+线扫·多配方", C_HW_FILL, C_HW),
        (2.6, "缺陷真值确认\n工艺/质检", C_PLC_FILL, C_PLC),
        (4.9, "标注 bbox\n检测+分类集", C_EDGE_FILL, C_EDGE),
        (7.2, "训练 det/cls\nTensorRT 部署", C_ALGO_FILL, C_ALGO),
        (9.5, "试运行回流\n漏检/误检补标", C_ACCENT, "#F5D5C0"),
    ]
    for i, (x, t, face, edge) in enumerate(items):
        _box(ax, x, 2.0, 2.1, 1.4, t, face if i < 4 else "#F5D5C0", edge if i < 4 else C_ACCENT, 8.5)
        if i < len(items) - 1:
            _arrow(ax, x + 2.1, 2.7, items[i + 1][0], 2.7)

    # feedback loop arrow
    ax.annotate(
        "",
        xy=(1.35, 2.0),
        xytext=(10.55, 2.0),
        arrowprops=dict(arrowstyle="-|>", color=C_ACCENT, lw=1.5, connectionstyle="arc3,rad=-0.35"),
    )
    ax.text(6.0, 0.7, "难例回流闭环：持续提升检出率与稳定性（必要支撑，非可选）", ha="center", fontsize=9, color=C_ACCENT, fontweight="bold")

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def set_run_font(run, name="宋体", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        from docx.shared import RGBColor

        run.font.color.rgb = color


def add_heading_cn(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        set_run_font(run, name="黑体", size=16 if level == 1 else 14 if level == 2 else 12, bold=True)
    return h


def add_para(doc, text, size=11, bold=False, first_line_indent=True):
    p = doc.add_paragraph()
    if first_line_indent:
        p.paragraph_format.first_line_indent = Cm(0.74)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.35
    run = p.add_run(text)
    set_run_font(run, size=size, bold=bold)
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    run = p.add_run(text)
    set_run_font(run, size=9, bold=True)
    return p


def add_figure(doc, image_path: Path, width_inches=6.2):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))
    p.paragraph_format.space_after = Pt(2)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = "Table Grid"
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        run = cell.paragraphs[0].add_run(h)
        set_run_font(run, name="黑体", size=10, bold=True)
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            run = cell.paragraphs[0].add_run(str(val))
            set_run_font(run, size=10)
    doc.add_paragraph()
    return table


def draw_area_recipe_flow(path: Path):
    """Area-scan cold path → recipe / exposure / thresholds."""
    _setup_cn_font()
    fig, ax = plt.subplots(figsize=(11.5, 5.0), dpi=160)
    ax.set_xlim(0, 12)
    ax.set_ylim(0, 5.2)
    ax.axis("off")
    fig.patch.set_facecolor("white")
    _subtitle(ax, "图7  面扫识别 → 配方 / 曝光 / 阈值适配（冷路径）")

    nodes = [
        (0.2, "进板光电\n触发面扫", C_PLC_FILL, C_PLC),
        (2.3, "2× 面扫落盘\n板面+对照面", C_HW_FILL, C_HW),
        (4.4, "材质·花纹·色系\n识别", "#FFF3E0", C_ACCENT),
        (6.5, "Recipe 决议\n(+MES 冲突策略)", C_EDGE_FILL, C_EDGE),
        (8.6, "下发曝光增益\n装载 model/阈值", C_EDGE_FILL, C_EDGE),
        (10.2, "线扫热路径\nInferRequest", C_ALGO_FILL, C_ALGO),
    ]
    y, w, h = 2.6, 1.85, 1.35
    for i, (x, t, face, edge) in enumerate(nodes):
        _box(ax, x, y, w if i < 5 else 1.6, h, t, face, edge, 7.5)
        if i < len(nodes) - 1:
            nx = nodes[i + 1][0]
            _arrow(ax, x + (w if i < 5 else 1.6), y + h / 2, nx, y + h / 2)

    _box(
        ax,
        0.3,
        0.35,
        5.5,
        1.6,
        "适配参数包（按 recipe_id）\n· ExposureProfile：线扫/面扫曝光增益\n· ModelSet：det+cls（可多配方共享）\n· ThresholdProfile：物理宽长面积阈值",
        C_BG,
        "#888888",
        7.5,
    )
    _box(
        ax,
        6.1,
        0.35,
        5.5,
        1.6,
        "要点\n· 配方数量 N 可扩展，不锁 4 套；初版可兼容四分色种子\n· 面扫不做毫米级缺陷主检；仅服务 recipe 与成像/阈值适配\n· 须在首条线扫推理前 recipe_ready；预算 ≤200 ms（不含机械等待）",
        "#FFF8F0",
        C_ACCENT,
        7.5,
    )

    fig.tight_layout()
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)


def generate_figures() -> dict[str, Path]:
    FIG_DIR.mkdir(parents=True, exist_ok=True)
    paths = {
        "arch": FIG_DIR / "fig1_system_architecture.png",
        "boundary": FIG_DIR / "fig2_module_boundary.png",
        "acq": FIG_DIR / "fig3_acquisition_flow.png",
        "triple": FIG_DIR / "fig4_triple_camera.png",
        "pipe": FIG_DIR / "fig5_single_pipeline.png",
        "sample": FIG_DIR / "fig6_sample_workflow.png",
        "area": FIG_DIR / "fig7_area_recipe_flow.png",
    }
    draw_system_architecture(paths["arch"])
    draw_module_boundary(paths["boundary"])
    draw_acquisition_flow(paths["acq"])
    draw_triple_camera_flow(paths["triple"])
    draw_single_pipeline(paths["pipe"])
    draw_sample_workflow(paths["sample"])
    draw_area_recipe_flow(paths["area"])
    return paths


def main():
    figs = generate_figures()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.2)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(48)
    run = title.add_run("板材封边缺陷智能检测设备")
    set_run_font(run, name="黑体", size=22, bold=True)

    st = doc.add_paragraph()
    st.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = st.add_run("视觉算法架构说明书（V3）")
    set_run_font(run, name="黑体", size=20, bold=True)

    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_before = Pt(20)
    run = meta.add_run(
        "文档版本：V3.0\n用途：算法架构说明、与现有软件模块边界、面扫配方适配、样本支撑要求\n"
        "相对 V2：① 合入 2 台面扫材质/花纹/色系识别与配方适配；② 推理环境固定 Windows IoT\n"
        "其余缺陷热路径与三路策略等复用 V2 已确认结论"
    )
    set_run_font(run, size=11)

    doc.add_page_break()

    add_heading_cn(doc, "1 概述", 1)
    add_heading_cn(doc, "1.1 文档目的", 2)
    add_para(
        doc,
        "本说明书描述板材封边缺陷视觉算法 V3 的架构设计结果。V3 在 V2 基础上合入物料清单中的两台面扫相机"
        "（材质 / 花纹 / 色系识别 → 配方、曝光与判定阈值适配），并将边端推理运行环境明确为 Windows IoT。"
        "缺陷检测热路径（三路线扫独立检测 + 板尾 OR 汇总、YOLO11s+n、落盘读图等）沿用 V2 结论。",
    )
    add_heading_cn(doc, "1.2 相对 V2 的变更", 2)
    add_table(
        doc,
        ["项", "V2", "V3"],
        [
            ["面扫", "架构正文未合入 BOM 双面扫闭环", "2× 面扫冷路径：识别 → recipe / 曝光 / 阈值"],
            ["推理 OS", "优先 Ubuntu（可沿用原 OS）", "固定 Windows IoT，与边端同机部署"],
            ["配方", "四分色为主叙述", "配置化 N 套（不锁 4；初版可兼容四分色种子）"],
            ["线扫缺陷主检", "3 路上中下独立 + 板尾 OR", "同 V2"],
            ["传图方式", "落盘 + 文件路径；暂不共享内存", "同 V2"],
            ["中心/MES/PLC 协议", "不改造", "同 V2"],
        ],
    )
    add_heading_cn(doc, "1.3 设计定位", 2)
    add_para(
        doc,
        "在现有边端客户端 / 视觉中心 / PLC / MES 体系上升级边端 AI 能力；不改造 MES、贴标与出料中控协议。"
        "算法以边端同机 Windows IoT 推理服务部署；边端提交图像本机文件路径与配方，算法读盘后回传结果"
        "（现阶段不使用共享内存）。进板阶段由面扫完成色系/花纹感知并决议 recipe；线扫热路径各路独立判定、板尾 OR 汇总。",
    )
    add_heading_cn(doc, "1.4 设计目标", 2)
    add_table(
        doc,
        ["项", "设计结果"],
        [
            ["线速", "最高 60 m/min"],
            ["像素精度", "约 0.03 mm 量级（以现场标定为准）"],
            ["检出目标", "板级缺陷检出率 ≥99%（验收口径以合同为准）"],
            ["软件边界", "保留现有边端 UI、视觉中心、PLC、MES；替换边端「AI 推理」并增强面扫配方适配"],
            ["硬件策略", "沿用 BOM/现场原硬件（2 面扫 + 3 线扫等）；先功能复原，再评估降本"],
            ["推理环境", "Windows IoT（边端同机）"],
            ["异常策略", "算法服务不可用时边端回退原模型槽位；面扫失败按 MES/降级策略且可审计"],
        ],
    )

    add_heading_cn(doc, "2 总体架构", 1)
    add_heading_cn(doc, "2.1 架构说明", 2)
    add_para(
        doc,
        "系统分为电气触发层、成像层（2 面扫 + 3 线扫）、边端既有软件与配方决议层、算法 V3 层与产线上层。"
        "面扫服务于冷路径配方适配；线扫服务于缺陷主检。结果仍由边端按原协议下发 PLC、上报视觉中心与 MES。图1 给出端到端关系。",
    )
    add_figure(doc, figs["arch"], 6.4)
    add_caption(doc, "图1  系统总体架构（V3）")

    add_heading_cn(doc, "2.2 分层职责", 2)
    add_table(
        doc,
        ["层级", "模块", "职责"],
        [
            ["电气", "光电、编码器、PLC", "进板触发面扫；线扫硬触发；接收 NG 等"],
            ["成像-面扫", "2× 面扫（如 MV-CS004-11GC）", "饰面/对照面成像，色系花纹识别用"],
            ["成像-线扫", "3× 线扫上/中/下（如 MV-CL024-91GM）", "Mono8 条带，缺陷主检"],
            ["边端现有", "采图 + UI + 通讯", "取流落盘；按 recipe 写曝光；业务外发"],
            ["边端增强", "面扫识别 + 配方决议", "输出 panel/edge tone、pattern；决议 recipe_id"],
            ["边端新增", "视觉算法服务 V3（Windows IoT）", "面扫分类（可同服务）+ 线扫 det/cls/度量/判定"],
            ["产线上层", "视觉中心、MES", "汇总复检、业务数据与色系先验（协议不变）"],
        ],
    )
    add_para(
        doc,
        "部署要点：Windows IoT 边端同机；落盘路径读图；先原硬件复原；视觉中心不跑 V3 推理。",
        first_line_indent=True,
    )

    add_heading_cn(doc, "3 面扫材质花纹色系识别与配方适配", 1)
    add_para(
        doc,
        "物料清单配置 2 台面扫。进板后双视角识别板面与对照面（或封边可见面）的材质、花纹与色系，"
        "映射为 recipe_id，并下发线扫曝光/增益，装载对应模型组与物理判定阈值。面扫不做白皮书级毫米缺陷主检。"
        "详细设计见项目 doc 中《01.面扫材质色系识别与配方适配》与原始需求《【面扫】材质花纹色系识别与配方适配》。",
    )
    add_figure(doc, figs["area"], 6.4)
    add_caption(doc, "图7  面扫识别 → 配方 / 曝光 / 阈值适配（冷路径）")

    add_heading_cn(doc, "3.1 流程与时序", 2)
    add_table(
        doc,
        ["步骤", "内容"],
        [
            ["1", "进板触发 2 面扫取流并落盘"],
            ["2", "识别输出：tone/pattern/confidence（特征可扩展）"],
            ["3", "Recipe 决议（面扫优先 + MES 校验/冲突策略可配）"],
            ["4", "下发线扫（及可选面扫）曝光增益；装载 model / thresholds"],
            ["5", "置 recipe_ready 后进入三路线扫缺陷热路径"],
        ],
    )
    add_para(
        doc,
        "时序：须在本板首条线扫 InferRequest 前完成；识别+决议+写参目标 ≤200 ms（不含机械等待）。超时走降级并告警，禁止静默未知配方。",
    )

    add_heading_cn(doc, "3.2 配方集合（可扩展）", 2)
    add_para(
        doc,
        "配方数量 N 不锁死为 4。客户既有浅/深板 × 浅/深封边四分色为兼容起点；最终套数与细分（花纹、反光等）"
        "在算法迭代中按样本可分性、检出收益与维护成本确认。实现上配置表驱动，禁止仅硬编码四枚举分支。",
    )
    add_table(
        doc,
        ["绑定项", "说明"],
        [
            ["recipe_id", "唯一枢纽；任意已注册字符串"],
            ["ExposureProfile", "线扫/面扫 exposure_us、gain_db 等"],
            ["ModelSet", "det+cls 权重；多 recipe 可共享同一权重"],
            ["ThresholdProfile", "各类缺陷宽/长/面积等物理阈值"],
            ["match 规则", "tone/pattern/MES 关键字 + 优先级"],
        ],
    )

    add_heading_cn(doc, "3.3 与 MES 冲突", 2)
    add_table(
        doc,
        ["情况", "默认行为"],
        [
            ["面扫 conf 达标", "采用面扫映射 recipe"],
            ["面扫 conf 低且 MES 有效", "采用 MES，告警"],
            ["面扫与 MES 冲突", "默认定面扫+记审计；可配置停检"],
            ["双源皆无或面扫掉线", "fallback 或同批上一板；必告警"],
        ],
    )

    add_heading_cn(doc, "4 算法 V3 与其他软件模块边界", 1)
    add_para(
        doc,
        "分界原则同 V2：电气与采图不入缺陷 det；V3 读落盘图并回传结构化结果。相对 V2，边端增加面扫识别与配方决议职责划分（图2）。",
    )
    add_figure(doc, figs["boundary"], 6.4)
    add_caption(doc, "图2  算法 V3 与周边软件模块边界")

    add_heading_cn(doc, "4.1 边界总表", 2)
    add_table(
        doc,
        ["模块", "归属", "负责内容", "不负责内容"],
        [
            [
                "视觉算法服务 V3",
                "本方案新增",
                "面扫分类（冷路径，可同服务）；线扫按路径读图；预处理/切片；检测分类度量；单路判定；板尾 OR",
                "相机 SDK 主控/硬触发发送；PLC 点位；MES/贴标协议；边端 UI/加密狗",
            ],
            [
                "边端客户端 + 配方决议",
                "既有增强",
                "工程方案/ROI；图像落盘；面扫→recipe 与冲突策略；写曝光；下发路径与 thresholds；展示回退",
                "V3 内部推理实现细节",
            ],
            [
                "边端采图模块",
                "边端内",
                "面扫+线扫 SDK 取流；按配方写曝光；组帧落盘",
                "缺陷识别与业务最终判定",
            ],
            [
                "PLC/光电/编码器",
                "电气",
                "进板与硬触发；运动相关；接收 NG",
                "图像算法",
            ],
            [
                "视觉中心 / MES",
                "既有",
                "多机汇总复检、业务数据；MES 色系先验",
                "边端本地 V3 推理",
            ],
        ],
    )

    add_heading_cn(doc, "4.2 采图与硬触发边界", 2)
    add_para(
        doc,
        "线扫硬触发由「边端采图 + PLC」完成；V3 不打开相机、不发 TRIG。边端落盘后 V3 在 Windows IoT 上按本地文件路径读图（图3）。",
    )
    add_figure(doc, figs["acq"], 6.4)
    add_caption(doc, "图3  采图与硬触发协作流程")
    add_table(
        doc,
        ["环节", "负责方", "设计结果"],
        [
            ["曝光/增益/行频/触发模式", "边端采图模块（SDK）", "工程方案与 recipe 绑定参数下发"],
            ["线扫硬触发", "PLC → 相机 TRIG", "禁用软触发量产"],
            ["面扫进板触发", "光电/PLC → 面扫", "到位采图供色系识别"],
            ["取流与落盘", "边端采图/客户端", "写本机约定目录"],
            ["算法输入", "边端 → V3", "本机路径 + board/recipe/阈值/可选 area_meta"],
            ["帧接入", "算法 V3", "按路径读文件；Windows IoT 部署"],
        ],
    )

    add_heading_cn(doc, "4.3 接口边界摘要", 2)
    add_para(
        doc,
        "边端 → V3：板材 ID、机台号、recipe_id（任意已注册）、可选 recipe_source/area_meta、"
        "三路线扫图像本机路径（不含面扫帧作 det 输入）、ROI、标定、thresholds（推荐边端按配方展开）。",
        first_line_indent=False,
    )
    add_table(
        doc,
        ["字段（V3 → 边端）", "说明"],
        [
            ["board_ok", "板级 OK / NG"],
            ["defects[]", "类别、相机位、框、宽/长/面积（mm）、置信度"],
            ["counts", "各类计数"],
            ["engine", "标识 scanly-v3"],
        ],
    )
    add_para(doc, "边端再按原协议写入 PLC、上报视觉中心。V3 不直连 PLC / MES。未知 recipe_id 须错误或降级，禁止静默。")

    add_heading_cn(doc, "5 三路线扫相机策略", 1)
    add_para(
        doc,
        "上、中、下三路主责不同成像区域，业务上无需跨相机框级融合。每路独立完成检测与判定，"
        "板材离开后仅做 OR / 计数汇总（图4）。与 V2 相同。",
    )
    add_figure(doc, figs["triple"], 6.4)
    add_caption(doc, "图4  三路相机独立检测与板尾汇总")
    add_table(
        doc,
        ["项", "设计结果"],
        [
            ["运行方式", "上 / 中 / 下三路独立检测、独立判定"],
            ["板级结论", "任一路 NG → 板 NG；列表拼接；计数相加"],
            ["跨相机融合", "不做框级同步融合与跨路等待"],
        ],
    )
    add_table(
        doc,
        ["相机", "主检内容"],
        [
            ["上", "上棱胶缝、上表面划伤/压痕、残胶、波浪纹；端头过长/短带（上缘可见）"],
            ["中", "侧面低于/高于饰面、开胶、封边带损伤、崩缺"],
            ["下", "下棱胶缝、下表面伤、残胶；端头过长/短带（下缘可见）"],
        ],
    )

    add_heading_cn(doc, "6 单路算法流程与模型", 1)
    add_para(
        doc,
        "线扫单路流水线同 V2：读盘 → 预处理 → ROI/门控 → 切片 → YOLO11s 检测 → YOLO11n-cls → 物理度量 → 阈值判定（图5）。"
        "配方切换时按 recipe_id 加载 model/阈值；当前仅驻留一组权重。",
    )
    add_figure(doc, figs["pipe"], 6.5)
    add_caption(doc, "图5  单路算法处理流水线")

    add_heading_cn(doc, "6.1 步骤说明", 2)
    add_table(
        doc,
        ["步骤", "内容", "说明"],
        [
            ["① 读本地文件", "按路径打开已落盘线扫图像", "不控相机；暂不共享内存"],
            ["② 预处理", "CLAHE / 归一化等 GPU 处理", "减轻后续网络负担"],
            ["③ ROI/门控", "封边区域 ROI；空窗跳过", "降低无效推理"],
            ["④ 切片", "1024 窗口，重叠可配", "应对极小目标"],
            ["⑤ 检测", "YOLO11s，输入默认 640", "高召回定位"],
            ["⑥ 分类", "YOLO11n-cls，裁剪约 224", "细分类与打假阳"],
            ["⑦ 度量", "像素→毫米", "使用标定 mm/px"],
            ["⑧ 判定", "可配置阈值（按 recipe）", "输出本路结果"],
        ],
    )
    add_heading_cn(doc, "6.2 模型配置", 2)
    add_table(
        doc,
        ["模型", "规格", "作用"],
        [
            ["检测", "YOLO11s，输入默认 640", "定位缺陷区域"],
            ["分类", "YOLO11n-cls，裁剪约 224", "细分类与误报抑制"],
            ["加速", "TensorRT FP16（Windows IoT）", "边端实时推理"],
            ["配方", "配置化 N 套 det+cls（可共享权重）", "按 recipe_id 热加载；套数迭代确认"],
            ["面扫分类", "规则特征或轻量 cls", "冷路径；不入线扫条带周期"],
        ],
    )

    add_heading_cn(doc, "7 缺陷判定（默认阈值，可配置）", 1)
    add_table(
        doc,
        ["缺陷", "默认条件（mm）"],
        [
            ["胶缝", "宽 >0.2 且 长 >5"],
            ["侧面低于饰面", "高 >0.3 且 长 >5"],
            ["短带", "长 >1"],
            ["过长", "长 >1"],
            ["划伤/压痕", "宽 >0.4 且 长 >5"],
            ["残胶", "宽 >1 且 长 >3"],
            ["开胶", "宽 >1 且 长 >5（可配置为检出即 NG）"],
            ["崩缺", "宽 >1 且 长 >1"],
            ["封边带损伤", "宽 >1 且 长 >1"],
            ["波浪纹", "宽 >0.7"],
        ],
    )
    add_para(doc, "全局默认表可按 recipe 叠加偏置；阈值由工程方案/配方下发，不写死在模型内。")

    add_heading_cn(doc, "8 样本收集与标注支撑要求", 1)
    add_para(
        doc,
        "V3 检出率与稳定性依赖现场真实样本。除线扫缺陷标注外，面扫双视角样图与当前启用配方全集均需覆盖（图6）。",
    )
    add_figure(doc, figs["sample"], 6.4)
    add_caption(doc, "图6  样本收集 · 标注 · 训练闭环")

    add_heading_cn(doc, "8.1 样本范围与规模", 2)
    add_table(
        doc,
        ["类别", "要求"],
        [
            ["缺陷类型", "覆盖约定缺陷：胶缝、侧面高低、短带/过长、划伤、残胶、开胶、崩缺、封边带损伤、波浪纹等"],
            ["色系配方", "当前启用配方全集代表性样本（N 不限 4）；面扫双视角同步留存"],
            ["相机视角", "上/中/下线扫 + 两面扫均可区分"],
            ["规模建议", "每类有效缺陷框建议 ≥1000 量级（极明显类可酌减）"],
            ["负样本/难例", "正常板面与易误报纹理；试运行漏检误检回流"],
        ],
    )
    add_heading_cn(doc, "8.2 标注与协作分工", 2)
    add_table(
        doc,
        ["事项", "甲方（客户/现场）", "乙方（算法方）"],
        [
            ["产线采图", "提供采图条件或授权存图，覆盖配方与班次", "提出采集清单与质量要求"],
            ["缺陷真值", "工艺/质检确认；协助难例收集", "类别体系与标注规范"],
            ["标注执行", "安排或委托标注；参与抽检与口径确认", "工具/规范、质检与清洗"],
            ["迭代", "持续提供漏检误检回流", "补训与版本发布；确认配方 N"],
            ["验收数据", "确认验收集代表性", "在约定数据集报告指标"],
        ],
    )

    add_heading_cn(doc, "9 硬件配置与性能", 1)
    add_heading_cn(doc, "9.1 实施策略", 2)
    add_para(
        doc,
        "优先沿用 BOM / 现场既有硬件（含 2 面扫 + 3 线扫、2080 Ti 等）完成功能复原与联调；"
        "功能稳定后再评估降本。不以换卡为功能上线前置。推理服务部署于 Windows IoT。",
    )
    add_heading_cn(doc, "9.2 硬件与运行环境基线", 2)
    add_table(
        doc,
        ["类别", "配置（原方案/现场常用）", "说明"],
        [
            ["面扫相机", "2× 如 MV-CS004-11GC + 配套镜头", "色系/花纹识别，非缺陷主检"],
            ["线扫相机", "3× 海康 MV-CL024-91GM（上/中/下）", "GigE · 硬触发 · Mono8"],
            ["镜头/光源", "工业镜头 + 光源及控制器", "按既有机位；配方改曝光增益"],
            ["触发链路", "编码器 + 光电 → PLC → 相机", "面扫进板触发 + 线扫硬触发"],
            ["边端工控机", "i7 级 CPU，16GB 起，SSD 2TB 级", "如原 HC 系工控机等"],
            ["边端 GPU", "RTX 2080 Ti（现场/原方案常用）", "现阶段主用"],
            ["算法 OS", "Windows IoT", "V3 固定；边端同机部署与运维"],
            ["图像传递", "本机落盘 + 文件绝对路径", "现阶段不使用共享内存"],
            ["视觉中心", "汇总/复检主机", "不运行 V3 推理"],
            ["网络/通讯", "工业千兆；PLC Modbus/S7；MES HTTP", "协议保持不变"],
        ],
    )
    add_heading_cn(doc, "9.3 性能目标", 2)
    add_table(
        doc,
        ["项", "目标"],
        [
            ["面扫冷路径", "识别+决议+写参 ≤200 ms（不含机械等待）"],
            ["单路线扫算法（不含采图/盘 IO）", "≤23 ms 量级（2080 Ti 裕度更宽）"],
            ["线速目标", "最高 60 m/min（可分档）"],
            ["视觉中心", "不运行本算法推理"],
        ],
    )

    add_heading_cn(doc, "10 实施阶段", 1)
    add_table(
        doc,
        ["阶段", "交付内容"],
        [
            [
                "一阶段（原硬件复原）",
                "Windows IoT 上单路线扫 det+cls+度量；边端对接与回退；手动/MES 选 recipe（可用四分色兼容种子）",
            ],
            [
                "二阶段（三路与配方表）",
                "三路独立并行；板尾 OR；配置化阈值/模型绑定；与视觉中心对齐",
            ],
            [
                "二点五阶段（面扫自适应）",
                "双面扫识别 → 自动 recipe/曝光/阈值；冲突策略；支持扩展 N 套",
            ],
            [
                "三阶段（达标与迭代确认）",
                "线速验收、难例回流；确认最终配方粒度/套数；可选硬件降本评估",
            ],
        ],
    )

    add_heading_cn(doc, "11 关联文档", 1)
    add_table(
        doc,
        ["类型", "路径或名称"],
        [
            ["原始需求", "scanly/doc/01-or/【面扫】材质花纹色系识别与配方适配.md"],
            ["面扫设计", "scanly/doc/02-dr/视觉算法/01.面扫材质色系识别与配方适配.md"],
            ["内部完整架构（V2 基线）", "scanly/doc/02-dr/视觉算法/视觉算法架构V2.md"],
            ["测试设计", "scanly/doc/03-tr/视觉算法/面扫材质色系识别与配方适配测试设计.md"],
            ["硬件 BOM", "scanly/doc/00-客户资料/鲲鹭-木板封边检测物料清单.xlsx"],
            ["生成脚本", "scanly/defects/tools/gen_arch_docx_v3.py"],
        ],
    )

    add_heading_cn(doc, "12 附则", 1)
    add_para(
        doc,
        "本说明书描述 V3 已确认的设计结果。除第 1.2 节所列相对 V2 的变更外，缺陷热路径、三路策略与性能预算等与 V2 保持一致。"
        "接口字段细节与内部实现以双方联调约定及后续技术附件为准。若硬件、软件版本或缺陷口径变更，应同步评估对架构边界与样本计划的影响。",
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT))
    print(f"Wrote figures under: {FIG_DIR}")
    print(f"Wrote: {OUT}")


if __name__ == "__main__":
    main()
